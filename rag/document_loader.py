from abc import ABC, abstractmethod
from dataclasses import dataclass

import pypdf
from docx import Document as DocxDocument
from docx.table import Table

from utils.file_utils import get_file_extension
from utils.helpers import clean_text


@dataclass
class Document:
    content: str
    metadata: dict
    source: str


class DocumentLoader(ABC):
    @abstractmethod
    def load(self, file_path: str) -> list[Document]:
        pass

    @abstractmethod
    def supports(self, file_path: str) -> bool:
        pass


class PDFLoader(DocumentLoader):
    def supports(self, file_path: str) -> bool:
        return get_file_extension(file_path) == ".pdf"

    def load(self, file_path: str) -> list[Document]:
        documents = []
        try:
            with open(file_path, "rb") as f:
                reader = pypdf.PdfReader(f)
                for page_num, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        text = clean_text(text)
                        documents.append(
                            Document(
                                content=text,
                                metadata={"page": page_num + 1, "source": file_path},
                                source=file_path,
                            )
                        )
        except Exception as e:
            raise ValueError(f"Failed to load PDF: {e!s}")
        return documents


class DocxLoader(DocumentLoader):
    def supports(self, file_path: str) -> bool:
        return get_file_extension(file_path) == ".docx"

    def _extract_table_text(self, table: Table) -> str:
        rows_text = []
        for row in table.rows:
            cell_texts = [cell.text.strip() for cell in row.cells]
            rows_text.append(" | ".join(cell_texts))
        return "\n".join(rows_text)

    def load(self, file_path: str) -> list[Document]:
        documents = []
        try:
            doc = DocxDocument(file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)

            for table in doc.tables:
                table_text = self._extract_table_text(table)
                if table_text.strip():
                    full_text.append(table_text)

            text = clean_text("\n".join(full_text))
            if text:
                documents.append(
                    Document(content=text, metadata={"source": file_path}, source=file_path)
                )
        except Exception as e:
            raise ValueError(f"Failed to load DOCX: {e!s}")
        return documents


class TextLoader(DocumentLoader):
    def supports(self, file_path: str) -> bool:
        return get_file_extension(file_path) == ".txt"

    def load(self, file_path: str) -> list[Document]:
        documents = []
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = clean_text(f.read())
                if text:
                    documents.append(
                        Document(
                            content=text,
                            metadata={"source": file_path},
                            source=file_path,
                        )
                    )
        except UnicodeDecodeError:
            try:
                with open(file_path, "r", encoding="latin-1") as f:
                    text = clean_text(f.read())
                    if text:
                        documents.append(
                            Document(
                                content=text,
                                metadata={"source": file_path},
                                source=file_path,
                            )
                        )
            except Exception as e:
                raise ValueError(f"Failed to load text file: {e!s}")
        except Exception as e:
            raise ValueError(f"Failed to load text file: {e!s}")
        return documents


class DocumentLoaderFactory:
    _loaders: list[DocumentLoader] | None = None

    @classmethod
    def _ensure_loaders(cls):
        if cls._loaders is None:
            cls._loaders = [PDFLoader(), DocxLoader(), TextLoader()]

    @classmethod
    def get_loader(cls, file_path: str) -> DocumentLoader | None:
        cls._ensure_loaders()
        assert cls._loaders is not None
        for loader in cls._loaders:
            if loader.supports(file_path):
                return loader
        return None

    @classmethod
    def load_document(cls, file_path: str) -> list[Document]:
        loader = cls.get_loader(file_path)
        if loader is None:
            raise ValueError(f"Unsupported file type: {file_path}")
        return loader.load(file_path)

    @classmethod
    def register_loader(cls, loader: DocumentLoader) -> None:
        cls._ensure_loaders()
        assert cls._loaders is not None
        cls._loaders.insert(0, loader)


def load_document(file_path: str) -> list[Document]:
    return DocumentLoaderFactory.load_document(file_path)


def load_documents(file_paths: list[str]) -> list[Document]:
    all_docs = []
    for path in file_paths:
        all_docs.extend(load_document(path))
    return all_docs
